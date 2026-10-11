#!/usr/bin/env python3
"""7A diagnostika to‘plamini savollar/*.json asosida qayta yasaydi.

Yasaladi:
  03_oquvchi_sorovnomasi.md, 04_kirish_testi_A.md, 05_yakuniy_test_B.md, 06_javoblar_kaliti.md
  jadvallar/7A_monitoring.xlsx
  chop_etish/*.pdf va chop_etish/7A_toliq_qollanma.docx

Talab: Python 3.9+, openpyxl; PDF/DOCX uchun pandoc va Chromium (bo‘lmasa o‘tkazib yuboriladi).

Ishlatish:
  python3 tools/build.py                         # hammasi
  python3 tools/build.py --pdfsiz                # faqat md + xlsx
  python3 tools/build.py --ismlar ismlar.txt --excel D:/7A_shaxsiy.xlsx
      # ism-familiyalar bilan to‘ldirilgan SHAXSIY Excel (repoga qo‘ymang!)
"""
import argparse
import html
import json
import os
import shutil
import subprocess
import tempfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as col
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
SAV = ROOT / "savollar"
HARF = "ABCD"
MAX_OQUVCHI = 35
BIRINCHI_QATOR = 6  # Excel’da 1-o‘quvchi qatori (A_kirish, B_yakuniy, Amaliy, Solishtirish)
DARAJALAR = [("Yuqori", 85), ("Yaxshi", 65), ("O‘rta", 45), ("Boshlang‘ich", 0)]


def load(name):
    return json.loads((SAV / name).read_text("utf-8"))


def bolim_guruhlari(test):
    """Bo‘limlar ketma-ket kelishini tekshiradi: [(bolim, [savollar]), ...]."""
    tartib = [b["id"] for b in test["bolimlar"]]
    guruhlar = []
    for q in test["savollar"]:
        if not guruhlar or guruhlar[-1][0] != q["bolim"]:
            guruhlar.append((q["bolim"], []))
        guruhlar[-1][1].append(q)
    assert [g[0] for g in guruhlar] == tartib, "Savollar bo‘limlar tartibida kelishi kerak"
    return guruhlar


# ---------------------------------------------------------------- Markdown

def test_md(test, raqam):
    bolim_nomi = {b["id"]: b for b in test["bolimlar"]}
    q_soni = len(test["savollar"])
    out = [
        f"# {raqam}. {test['nomi']}",
        "",
        "| Sinf | Fan | Vaqt | Savollar | Qachon |",
        "|------|-----|------|----------|--------|",
        f"| {test['sinf']} | {test['fan']} | {test['vaqt_daqiqa']} daqiqa | {q_soni} ta | {test['qachon']} |",
        "",
        "**F.I.:** ______________________________ **Sana:** ____________ **Kompyuter №:** ______",
        "",
        f"> **Ko‘rsatma.** {test['korsatma']}",
        "",
        "> *Ustoz uchun:* saytda — `…/7a/` sahifasi; qog‘ozda — `chop_etish/` dagi PDF va `javoblar_varaqasi.pdf`. "
        "Kalit — `06_javoblar_kaliti.md`.",
        "",
    ]
    n = 0
    for i, (bid, savollar) in enumerate(bolim_guruhlari(test), 1):
        b = bolim_nomi[bid]
        out += [f"## {i}-bo‘lim. {b['nomi']}", ""]
        for q in savollar:
            n += 1
            out.append(f"**{n}.** {q['matn']}")
            out.append("")
            for j, v in enumerate(q["variantlar"]):
                out.append(f"- **{HARF[j]})** {v}")
            out.append("")
    return "\n".join(out)


XATO_TAHLILI = [
    (3, "«Operativ xotira» va «qattiq disk» farqini bilmaydi — kompyuter tuzilishini takrorlash kerak."),
    (8, "«Probel ko‘p bosish» yoki «shriftni kattalashtirish» tanlansa — Word’da tekislash odati yo‘q, matnni bo‘sh joy bilan «suradi»."),
    (10, "Backspace/Delete adashtirilsa — klaviatura bilan kam ishlagan."),
    (13, "Yig‘indi (20 yoki 13) tanlansa — takrorlashni qo‘shish deb tushunadi; «1010» yoki «58» — sonlarni yonma-yon yozgan, sikl ma’nosini bilmaydi."),
    (15, "Oraliq natija (7 yoki 15) tanlansa — algoritmni oxirigacha kuzatmagan; o‘zgaruvchi tushunchasi zaif."),
    (18, "1000 tanlansa — o‘nlik «kilo» bilan adashtiradi (informatikada 1024)."),
    (20, "Bitni baytga aylantirmasdan taqqoslagan — birliklarni o‘tkazishni mashq qilish kerak."),
    (22, "«101» yoki «110» (sonning o‘zi) tanlansa — ikkilik yozuv qiymatini tushunmaydi."),
    (24, "«12₂», «102₂» yoki «4₂/6₂» tanlansa — ikkilik sistemada o‘nlikdagidek qo‘shgan; ikkilikda 2 raqami yo‘qligini bilmaydi."),
    (27, "Sayt manzili, e-pochta va telefon raqamini adashtiradi — Internet manzillari mavzusiga e’tibor."),
    (29, "Firibgar xabariga «ishonadi» — XAVF belgisi! Bu o‘quvchi bilan alohida gaplashing va ota-onasiga xabar bering."),
]


def kalit_md(a, b):
    bolim_nomi = {x["id"]: x["nomi"] for x in a["bolimlar"]}
    qiyin = {1: "oson", 2: "o‘rta", 3: "qiyin"}
    out = [
        "# 06. Javoblar kaliti (faqat ustoz uchun)",
        "",
        "> ⚠️ Bu faylni o‘quvchilarga ko‘rsatmang va chop etilgan nusxasini sinfda qoldirmang.",
        "> Saytda test topshirilganda javoblar server tomonda tekshiriladi — o‘quvchi kalitni ko‘rmaydi.",
        "",
        "## 1. Qisqa kalit",
        "",
        "| № | A variant | B variant | Bo‘lim | Tekshiriladigan ko‘nikma | Qiyinlik |",
        "|--:|:--:|:--:|--------|--------------------------|----------|",
    ]
    for i, (qa, qb) in enumerate(zip(a["savollar"], b["savollar"]), 1):
        out.append(
            f"| {i} | **{HARF[qa['javob']]}** | **{HARF[qb['javob']]}** | {qa['bolim']} | "
            f"{qa['kunikma']} / {qb['kunikma']} | {qiyin[qa['qiyinlik']]} |"
        )
    satr_a = " ".join(f"{i}{HARF[q['javob']]}" for i, q in enumerate(a["savollar"], 1))
    satr_b = " ".join(f"{i}{HARF[q['javob']]}" for i, q in enumerate(b["savollar"], 1))
    out += [
        "",
        "**Bir qatorda (tez tekshirish uchun):**",
        "",
        f"- A: `{satr_a}`",
        f"- B: `{satr_b}`",
        "",
        "## 2. Bo‘limlar xaritasi",
        "",
        "| Bo‘lim | Mavzu | Qism | Manba | Savollar |",
        "|--------|-------|------|-------|----------|",
    ]
    n = 0
    for bid, savollar in bolim_guruhlari(a):
        info = next(x for x in a["bolimlar"] if x["id"] == bid)
        out.append(
            f"| {bid} | {bolim_nomi[bid]} | {info['qism']} | {info['manba']} | {n + 1}–{n + len(savollar)} |"
        )
        n += len(savollar)
    out += [
        "",
        "**Tayanch qism** = 1–15-savollar (7-sinfga tayyorlik). **7-sinf qismi** = 16–30-savollar.",
        "B variantdagi har bir savol shu raqamli A-savol bilan **bir xil ko‘nikmani, bir xil qiyinlikda** tekshiradi.",
        "",
        "## 3. Xato javoblar nimani anglatadi (diagnostika)",
        "",
        "Noto‘g‘ri variantlar tasodifiy emas — har biri tipik tushunmovchilikni ko‘rsatadi. "
        "Saytdagi admin panelida har bir o‘quvchi qaysi variantni tanlaganini ko‘rasiz.",
        "",
        "| Savol № | Agar o‘quvchi noto‘g‘ri javob bersa … |",
        "|:------:|----------------------------------------|",
    ]
    for i, matn in XATO_TAHLILI:
        out.append(f"| {i} | {matn} |")
    for nomi, test in (("A", a), ("B", b)):
        out += ["", f"## {4 if nomi == 'A' else 5}. {nomi} variant — to‘g‘ri javoblar matni bilan", ""]
        for i, q in enumerate(test["savollar"], 1):
            out.append(f"{i}. {q['matn']} → **{HARF[q['javob']]}) {q['variantlar'][q['javob']]}**")
    out.append("")
    return "\n".join(out)


def sorovnoma_qism_md(s, sarlavha_daraja="##"):
    out = [
        f"{sarlavha_daraja} {s['nomi']}",
        "",
        f"*Vaqt: ~{s['vaqt_daqiqa']} daqiqa.* " + ("**Anonim — ism yozilmaydi.**" if s.get("anonim") else
                                                     "**F.I.:** ______________________________ **Sana:** ____________"),
        "",
        f"> {s['korsatma']}",
        "",
    ]
    for b in s["bolimlar"]:
        out += [f"{sarlavha_daraja}# {b['nomi']}", ""]
        for q in b["savollar"]:
            tur = q["tur"]
            qo = ""
            if tur == "multi":
                qo = f" *(bir nechtasini belgilang{', ko‘pi bilan ' + str(q['max']) + ' ta' if q.get('max') else ''})*"
            elif tur == "single":
                qo = " *(bittasini belgilang)*"
            out.append(f"**{q['matn']}**{qo}")
            out.append("")
            if tur in ("single", "multi"):
                variantlar = list(q["variantlar"]) + (["Boshqa: ____________________"] if q.get("boshqa") else [])
                out.append(" &nbsp; ".join(f"☐&nbsp;{v}" for v in variantlar))
            elif tur == "scale":
                kat = " &nbsp; ".join(f"☐&nbsp;{i}" for i in range(q["min"], q["max"] + 1))
                out.append(f"{q['belgilar'][0]} — {kat} — {q['belgilar'][1]}")
            elif tur == "matrix":
                out.append("| | " + " | ".join(q["ustunlar"]) + " |")
                out.append("|---|" + "|".join(":-:" for _ in q["ustunlar"]) + "|")
                for r in q["qatorlar"]:
                    out.append(f"| {r} | " + " | ".join("☐" for _ in q["ustunlar"]) + " |")
            elif tur == "text":
                out.append("_______________________________________________________________")
            elif tur == "textarea":
                out.append("\n\n".join(["_______________________________________________________________"] * 3))
            out.append("")
    return "\n".join(out)


def sorovnoma_md(s, f):
    return "\n".join([
        "# 03. O‘quvchi so‘rovnomalari",
        "",
        "Bu faylda ikkita so‘rovnoma bor. Ikkalasi ham saytda (`…/7a/`) onlayn mavjud — qog‘oz varianti internet bo‘lmasa kerak.",
        "",
        "| So‘rovnoma | Qachon | Kim | Nima uchun |",
        "|-----------|--------|-----|------------|",
        "| **A. Tanishuv so‘rovnomasi** | Oktabr, 1-hafta | Ismli | O‘quvchini tanish: qiziqishlar, kompyuter imkoniyati, o‘qish uslubi, o‘z-o‘zini baholash |",
        "| **B. Fikr so‘rovnomasi** | Yanvar (B test bilan) | **Anonim** | Darslaringiz haqida halol fikr + fanga munosabat va o‘z-o‘zini baholashni oktabr bilan solishtirish |",
        "",
        "**O‘tkazish bo‘yicha maslahatlar:**",
        "",
        "- So‘rovnomani testdan **oldin** bering — tanishuv kayfiyatini yaratadi.",
        "- Ovoz chiqarib o‘qing: *«To‘g‘ri yoki noto‘g‘ri javob yo‘q. Javoblarni faqat men ko‘raman.»*",
        "- Sekin yozadiganlarga ochiq savollarni qisqa yozishga ruxsat bering — asosiysi belgilanadigan savollar.",
        "- Anonim so‘rovnomada qog‘ozlarni **yopiq qutiga** yig‘ing, ismlar so‘ralmaydi.",
        "",
        "---",
        "",
        sorovnoma_qism_md(s).replace(f"## {s['nomi']}", f"## A. {s['nomi']}", 1),
        "---",
        "",
        sorovnoma_qism_md(f).replace(f"## {f['nomi']}", f"## B. {f['nomi']}", 1),
    ])


# ---------------------------------------------------------------- Excel

INGICHKA = Side(style="thin", color="999999")
CHEGARA = Border(left=INGICHKA, right=INGICHKA, top=INGICHKA, bottom=INGICHKA)
SARLAVHA_FILL = PatternFill("solid", fgColor="1F3A5F")
SARLAVHA_FONT = Font(bold=True, color="FFFFFF")
KALIT_FILL = PatternFill("solid", fgColor="FFF2CC")
HISOB_FILL = PatternFill("solid", fgColor="EEF3FA")
RANG = {
    "Yuqori": "C6EFCE", "Yaxshi": "BDD7EE", "O‘rta": "FFEB9C", "Boshlang‘ich": "FFC7CE",
    "Mos": "C6EFCE", "Qisman mos": "FFEB9C", "Mos emas": "FFC7CE",
    "Ilg‘or": "C6EFCE", "Asosiy": "BDD7EE", "Yordam": "FFC7CE",
    "↑": "C6EFCE", "=": "EDEDED", "↓": "FFC7CE",
}


def sarlavha(ws, qator, ustunlar):
    for c in ustunlar:
        cell = ws.cell(qator, c)
        cell.fill, cell.font, cell.border = SARLAVHA_FILL, SARLAVHA_FONT, CHEGARA
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def rang_qoidalari(ws, diapazon, qiymatlar):
    for q in qiymatlar:
        ws.conditional_formatting.add(
            diapazon, CellIsRule(operator="equal", formula=[f'"{q}"'], fill=PatternFill("solid", fgColor=RANG[q])))


def oxirgi():
    return BIRINCHI_QATOR + MAX_OQUVCHI - 1


def daraja_formula(x):
    return (f'IF({x}="","",IF({x}>=Darajalar!$B$2,"Yuqori",IF({x}>=Darajalar!$B$3,"Yaxshi",'
            f'IF({x}>=Darajalar!$B$4,"O‘rta","Boshlang‘ich"))))')


def test_varagi(wb, nomi, test, amaliy_ustun):
    ws = wb.create_sheet(nomi)
    qs = test["savollar"]
    n = len(qs)
    q0 = 3  # C ustun
    guruhlar = bolim_guruhlari(test)
    s0 = q0 + n  # bo‘limlar boshlanadigan ustun
    c_tayanch, c_yetti, c_ball, c_foiz = s0 + len(guruhlar), s0 + len(guruhlar) + 1, s0 + len(guruhlar) + 2, s0 + len(guruhlar) + 3
    c_daraja, c_10, c_mos, c_guruh = c_foiz + 1, c_foiz + 2, c_foiz + 3, c_foiz + 4
    oxir = oxirgi()

    ws["A1"] = f"7A — {test['nomi']}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = ("Har bir o‘quvchi qatoriga u belgilagan harfni (A/B/C/D) yozing. Bo‘sh katak = javob bermagan. "
                "Sariq qatordagi kalitni o‘zgartirmang. Kulrang ustunlar avtomatik hisoblanadi.")
    ws["A2"].alignment = Alignment(wrap_text=False)

    # 3-qator: bo‘lim belgilari, 4-qator: sarlavhalar, 5-qator: kalit
    c = q0
    for bid, savollar in guruhlar:
        ws.cell(3, c, bid).alignment = Alignment(horizontal="center")
        ws.merge_cells(start_row=3, start_column=c, end_row=3, end_column=c + len(savollar) - 1)
        c += len(savollar)
    sarlavha(ws, 3, range(q0, q0 + n))
    ws.cell(4, 1, "№")
    ws.cell(4, 2, "F.I.Sh.")
    for i in range(n):
        ws.cell(4, q0 + i, i + 1)
    for k, (bid, _) in enumerate(guruhlar):
        ws.cell(4, s0 + k, bid)
    for c_, t in ((c_tayanch, "Tayanch %"), (c_yetti, "7-sinf %"), (c_ball, f"Ball ({n})"), (c_foiz, "Jami %"),
                  (c_daraja, "Daraja"), (c_10, "10 ball"), (c_mos, "7-sinfga moslik"), (c_guruh, "Guruh (taklif)")):
        ws.cell(4, c_, t)
    sarlavha(ws, 4, range(1, c_guruh + 1))
    ws.cell(5, 2, "KALIT →").font = Font(bold=True)
    for i, q in enumerate(qs):
        cell = ws.cell(5, q0 + i, HARF[q["javob"]])
        cell.fill, cell.font, cell.border = KALIT_FILL, Font(bold=True), CHEGARA
        cell.alignment = Alignment(horizontal="center")

    tayanch = [k for k, (bid, _) in enumerate(guruhlar) if next(b for b in test["bolimlar"] if b["id"] == bid)["qism"] == "tayanch"]
    yetti = [k for k in range(len(guruhlar)) if k not in tayanch]
    t_soni = sum(len(guruhlar[k][1]) for k in tayanch)
    y_soni = n - t_soni

    for r in range(BIRINCHI_QATOR, oxir + 1):
        k = r - BIRINCHI_QATOR + 1
        ws.cell(r, 1, k)
        ws.cell(r, 2, f'=IF(Royxat!B{k + 1}="","",Royxat!B{k + 1})')
        javoblar = f"${col(q0)}{r}:${col(q0 + n - 1)}{r}"
        c = q0
        for j, (bid, savollar) in enumerate(guruhlar):
            a, b = col(c), col(c + len(savollar) - 1)
            ws.cell(r, s0 + j, f'=IF(COUNTA({javoblar})=0,"",SUMPRODUCT(--(UPPER(TRIM({a}{r}:{b}{r}))={a}$5:{b}$5)))')
            c += len(savollar)
        S = lambda j: f"{col(s0 + j)}{r}"
        ws.cell(r, c_tayanch, f'=IF({S(0)}="","",ROUND(({"+".join(S(j) for j in tayanch)})/{t_soni}*100,0))')
        ws.cell(r, c_yetti, f'=IF({S(0)}="","",ROUND(({"+".join(S(j) for j in yetti)})/{y_soni}*100,0))')
        ws.cell(r, c_ball, f'=IF({S(0)}="","",SUM({S(0)}:{S(len(guruhlar) - 1)}))')
        ws.cell(r, c_foiz, f'=IF({col(c_ball)}{r}="","",ROUND({col(c_ball)}{r}/{n}*100,0))')
        F = f"{col(c_foiz)}{r}"
        T = f"{col(c_tayanch)}{r}"
        ws.cell(r, c_daraja, "=" + daraja_formula(F))
        ws.cell(r, c_10, f'=IF({F}="","",ROUND({F}/10,0))')
        ws.cell(r, c_mos, f'=IF({T}="","",IF({T}>=Darajalar!$B$3,"Mos",IF({T}>=Darajalar!$B$4,"Qisman mos","Mos emas")))')
        am = f"Amaliy!{amaliy_ustun}{r}"
        ws.cell(r, c_guruh, f'=IF({F}="","",IF(AND({F}>=75,OR({am}="",{am}>=75)),"Ilg‘or",'
                            f'IF(OR({F}<45,AND({am}<>"",{am}<45)),"Yordam","Asosiy")))')
        for c_ in range(1, c_guruh + 1):
            ws.cell(r, c_).border = CHEGARA
            if c_ >= s0:
                ws.cell(r, c_).fill = HISOB_FILL
            if c_ >= q0:
                ws.cell(r, c_).alignment = Alignment(horizontal="center")

    # Sinf o‘rtachasi va savol bo‘yicha tahlil
    r_ort, r_sav = oxir + 2, oxir + 3
    ws.cell(r_ort, 2, "Sinf o‘rtacha").font = Font(bold=True)
    ws.cell(r_sav, 2, "Savol bo‘yicha to‘g‘ri %").font = Font(bold=True)
    for c_ in list(range(s0, c_foiz + 1)) + [c_10]:
        L = col(c_)
        ws.cell(r_ort, c_, f'=IFERROR(ROUND(AVERAGE({L}{BIRINCHI_QATOR}:{L}{oxir}),1),"")').font = Font(bold=True)
    foiz_diap = f"${col(c_foiz)}${BIRINCHI_QATOR}:${col(c_foiz)}${oxir}"
    for i in range(n):
        L = col(q0 + i)
        ws.cell(r_sav, q0 + i, f'=IF(COUNT({foiz_diap})=0,"",ROUND(SUMPRODUCT(--(UPPER(TRIM({L}${BIRINCHI_QATOR}:{L}${oxir}))={L}$5))/COUNT({foiz_diap})*100,0))')
        ws.cell(r_sav, q0 + i).alignment = Alignment(horizontal="center")
    ws.conditional_formatting.add(
        f"{col(q0)}{r_sav}:{col(q0 + n - 1)}{r_sav}",
        FormulaRule(formula=[f'AND({col(q0)}{r_sav}<>"",{col(q0)}{r_sav}<40)'], fill=PatternFill("solid", fgColor="FFC7CE")))
    ws.cell(r_sav + 1, 2, "40 % dan past (qizil) savollar — butun sinf uchun qiyin mavzu.").font = Font(italic=True, color="666666")

    # Javob kataklari: to‘g‘ri — yashil, noto‘g‘ri — qizil
    javob_diap = f"{col(q0)}{BIRINCHI_QATOR}:{col(q0 + n - 1)}{oxir}"
    tl = f"{col(q0)}{BIRINCHI_QATOR}"
    ws.conditional_formatting.add(javob_diap, FormulaRule(
        formula=[f'AND({tl}<>"",UPPER(TRIM({tl}))={col(q0)}$5)'], fill=PatternFill("solid", fgColor="C6EFCE")))
    ws.conditional_formatting.add(javob_diap, FormulaRule(
        formula=[f'AND({tl}<>"",UPPER(TRIM({tl}))<>{col(q0)}$5)'], fill=PatternFill("solid", fgColor="FFC7CE")))
    dv = DataValidation(type="list", formula1='"A,B,C,D"', allow_blank=True, showErrorMessage=True,
                        errorStyle="warning", errorTitle="Javob", error="Faqat A, B, C yoki D yozing.")
    ws.add_data_validation(dv)
    dv.add(javob_diap)

    rang_qoidalari(ws, f"{col(c_daraja)}{BIRINCHI_QATOR}:{col(c_daraja)}{oxir}", ["Yuqori", "Yaxshi", "O‘rta", "Boshlang‘ich"])
    rang_qoidalari(ws, f"{col(c_mos)}{BIRINCHI_QATOR}:{col(c_mos)}{oxir}", ["Mos", "Qisman mos", "Mos emas"])
    rang_qoidalari(ws, f"{col(c_guruh)}{BIRINCHI_QATOR}:{col(c_guruh)}{oxir}", ["Ilg‘or", "Asosiy", "Yordam"])

    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 30
    for i in range(n):
        ws.column_dimensions[col(q0 + i)].width = 4
    for c_ in range(s0, c_guruh + 1):
        ws.column_dimensions[col(c_)].width = 9
    for c_ in (c_daraja, c_mos, c_guruh):
        ws.column_dimensions[col(c_)].width = 14
    ws.row_dimensions[4].height = 30
    ws.freeze_panes = ws.cell(BIRINCHI_QATOR, q0)
    return {"tayanch": col(c_tayanch), "yetti": col(c_yetti), "foiz": col(c_foiz), "daraja": col(c_daraja),
            "mos": col(c_mos), "guruh": col(c_guruh), "bolimlar": [col(s0 + j) for j in range(len(guruhlar))],
            "bolim_idlar": [g[0] for g in guruhlar], "bolim_soni": [len(g[1]) for g in guruhlar]}


def build_excel(path, a, b, ismlar=None):
    wb = Workbook()
    yo = wb.active
    yo.title = "Yoriqnoma"
    matn = [
        ("7A — monitoring jadvali (Informatika)", Font(bold=True, size=16)),
        ("", None),
        ("Qanday ishlatish:", Font(bold=True)),
        ("1. «Royxat» varag‘iga o‘quvchilar ism-familiyasini yozing (boshqa varaqlarga o‘zi o‘tadi).", None),
        ("2. «A_kirish» — oktabrdagi kirish testi: har bir savolga o‘quvchi belgilagan harfni (A/B/C/D) yozing.", None),
        ("3. «Amaliy» — amaliy topshiriq: har topshiriqqa 0, 1 yoki 2 ball (oktabr va yanvar ustunlari).", None),
        ("4. «B_yakuniy» — yanvardagi yakuniy test (B variant): xuddi shunday.", None),
        ("5. «Solishtirish» — avtomatik: o‘sish (Δ), normallashtirilgan o‘sish (g), daraja o‘zgarishi, sinf xulosasi va diagrammalar.", None),
        ("6. «Xarita» — kompyuter xonasidagi o‘tirish xaritasi (ism yodlash uchun).", None),
        ("7. «Darajalar» — chegaralar (85 / 65 / 45). O‘zgartirsangiz, hamma joyda qayta hisoblanadi.", None),
        ("", None),
        ("Muhim:", Font(bold=True)),
        ("• Saytda test o‘tkazgan bo‘lsangiz: admin → «CSV» → Q01…Q30 ustunlarini nusxalab, «A_kirish» ning C6 katagidan joylashtiring.", None),
        ("• Qoidalar va talqin: 10_tahlil_va_darajalar.md.", None),
        ("• Bu fayl ism-familiyalar bilan to‘ldirilgach SHAXSIY hujjat — ochiq joyga (GitHub, umumiy guruh) yuklamang.", None),
        ("• Yanvargacha zaxira nusxasini saqlang (fleshka / Google Drive).", None),
    ]
    for i, (t, f) in enumerate(matn, 1):
        yo.cell(i, 1, t)
        if f:
            yo.cell(i, 1).font = f
    yo.column_dimensions["A"].width = 130

    # Royxat
    ro = wb.create_sheet("Royxat")
    ust = ["№", "F.I.Sh.", "Laqab (qanday chaqirish)", "Guruh (1/2)", "Kompyuter №", "Qiziqishi",
           "Sinf rahbari izohi", "Ota-ona javobi", "Kuzatuv (1–3)", "Izoh"]
    kenglik = [4, 34, 18, 10, 11, 22, 34, 24, 11, 34]
    for i, (t, w) in enumerate(zip(ust, kenglik), 1):
        ro.cell(1, i, t)
        ro.column_dimensions[col(i)].width = w
    sarlavha(ro, 1, range(1, len(ust) + 1))
    for k in range(1, MAX_OQUVCHI + 1):
        ro.cell(k + 1, 1, k)
        if ismlar and k <= len(ismlar):
            ro.cell(k + 1, 2, ismlar[k - 1])
        for c_ in range(1, len(ust) + 1):
            ro.cell(k + 1, c_).border = CHEGARA
    ro.freeze_panes = "C2"

    # Darajalar
    da = wb.create_sheet("Darajalar")
    da["A1"], da["B1"], da["C1"], da["D1"] = "Daraja", "Quyi chegara (%)", "10 ballik", "Ma’nosi"
    izoh = ["Mavzuni to‘liq biladi", "Asosiy narsalarni biladi, kichik bo‘shliqlar", "Yarmini biladi, yordam kerak",
            "Tayanch bilimlar yetishmaydi"]
    ball = ["9–10", "7–8", "5–6", "0–4"]
    for i, ((d, ch), iz, bl) in enumerate(zip(DARAJALAR, izoh, ball), 2):
        da.cell(i, 1, d)
        da.cell(i, 2, ch)
        da.cell(i, 3, bl)
        da.cell(i, 4, iz)
        da.cell(i, 1).fill = PatternFill("solid", fgColor=RANG[d])
    sarlavha(da, 1, range(1, 5))
    rows = [
        ("7-sinfga moslik (Tayanch % bo‘yicha):", None),
        ("Mos", "Tayanch % ≥ «Yaxshi» chegarasi (65)"),
        ("Qisman mos", "Tayanch % ≥ «O‘rta» chegarasi (45)"),
        ("Mos emas", "Tayanch % < 45"),
        ("", None),
        ("Normallashtirilgan o‘sish g = (B% − A%) / (100 − A%):", None),
        ("≥ 0,7", "Yuqori o‘sish"),
        ("0,3 – 0,69", "O‘rtacha o‘sish (normal)"),
        ("0 – 0,29", "Past o‘sish"),
        ("< 0", "Pasaygan — alohida tahlil kerak"),
    ]
    for i, (x, y) in enumerate(rows, 8):
        da.cell(i, 1, x).font = Font(bold=y is None)
        if y:
            da.cell(i, 2, y)
    for c_, w in zip("ABCD", (34, 40, 10, 44)):
        da.column_dimensions[c_].width = w

    # Amaliy (test varaqlaridan oldin — formulalar unga murojaat qiladi)
    am = wb.create_sheet("Amaliy")
    am["A1"] = "7A — Amaliy topshiriq (07_amaliy_topshiriq.md). Har topshiriq: 0, 1 yoki 2 ball. Bajarilmagan topshiriqni bo‘sh qoldiring."
    am["A1"].font = Font(bold=True)
    am.cell(3, 3, "1-bosqich (oktabr)")
    am.merge_cells("C3:J3")
    am.cell(3, 11, "2-bosqich (yanvar)")
    am.merge_cells("K3:R3")
    am.cell(3, 20, "Yozish tezligi (belgi/daq)")
    am.merge_cells("T3:V3")
    sarlavha(am, 3, list(range(3, 11)) + list(range(11, 19)) + [20, 21, 22])
    vazifa = ["1. Papka", "2. Tezlik", "3. Bezash", "4. Saqlash", "5. Scratch", "6. Internet"]
    hdr = ["№", "F.I.Sh."] + vazifa + ["Ball", "%"] + vazifa + ["Ball", "%", "O‘sish (%)", "Oktabr", "Yanvar", "O‘sish"]
    for i, t in enumerate(hdr, 1):
        am.cell(4, i, t)
    sarlavha(am, 4, range(1, len(hdr) + 1))
    oxir = oxirgi()
    for r in range(BIRINCHI_QATOR, oxir + 1):
        k = r - BIRINCHI_QATOR + 1
        am.cell(r, 1, k)
        am.cell(r, 2, f'=IF(Royxat!B{k + 1}="","",Royxat!B{k + 1})')
        am.cell(r, 9, f'=IF(COUNT(C{r}:H{r})=0,"",SUM(C{r}:H{r}))')
        am.cell(r, 10, f'=IF(I{r}="","",ROUND(I{r}/(COUNT(C{r}:H{r})*2)*100,0))')
        am.cell(r, 17, f'=IF(COUNT(K{r}:P{r})=0,"",SUM(K{r}:P{r}))')
        am.cell(r, 18, f'=IF(Q{r}="","",ROUND(Q{r}/(COUNT(K{r}:P{r})*2)*100,0))')
        am.cell(r, 19, f'=IF(OR(J{r}="",R{r}=""),"",R{r}-J{r})')
        am.cell(r, 22, f'=IF(OR(T{r}="",U{r}=""),"",U{r}-T{r})')
        for c_ in range(1, 23):
            am.cell(r, c_).border = CHEGARA
            if c_ >= 3:
                am.cell(r, c_).alignment = Alignment(horizontal="center")
            if c_ in (9, 10, 17, 18, 19, 22):
                am.cell(r, c_).fill = HISOB_FILL
    dv = DataValidation(type="whole", operator="between", formula1="0", formula2="2", allow_blank=True,
                        errorStyle="warning", error="Faqat 0, 1 yoki 2.")
    am.add_data_validation(dv)
    dv.add(f"C{BIRINCHI_QATOR}:H{oxir}")
    dv.add(f"K{BIRINCHI_QATOR}:P{oxir}")
    am.column_dimensions["A"].width = 4
    am.column_dimensions["B"].width = 30
    for c_ in range(3, 23):
        am.column_dimensions[col(c_)].width = 9
    am.row_dimensions[4].height = 30
    am.freeze_panes = "C6"

    A = test_varagi(wb, "A_kirish", a, "J")
    B = test_varagi(wb, "B_yakuniy", b, "R")

    # Solishtirish
    so = wb.create_sheet("Solishtirish")
    so["A1"] = "7A — Oktabr (A) ↔ Yanvar (B) solishtirish. Bu varaq to‘liq avtomatik (faqat NB soni va Izohni yozasiz)."
    so["A1"].font = Font(bold=True, size=12)
    hdr = ["№", "F.I.Sh.", "A: Jami %", "B: Jami %", "O‘sish Δ", "g", "Daraja (okt)", "Daraja (yan)",
           "O‘zgarish", "Δ talqini", "Tayanch A", "Tayanch B", "7-sinf A", "7-sinf B", "Amaliy okt",
           "Amaliy yan", "Guruh (okt)", "NB soni", "Izoh"]
    for i, t in enumerate(hdr, 1):
        so.cell(4, i, t)
    sarlavha(so, 4, range(1, len(hdr) + 1))
    D = "Darajalar!$A$2:$A$5"
    for r in range(BIRINCHI_QATOR, oxir + 1):
        k = r - BIRINCHI_QATOR + 1
        f = {
            1: k,
            2: f'=IF(Royxat!B{k + 1}="","",Royxat!B{k + 1})',
            3: f"=A_kirish!{A['foiz']}{r}",
            4: f"=B_yakuniy!{B['foiz']}{r}",
            5: f'=IF(OR(C{r}="",D{r}=""),"",D{r}-C{r})',
            6: f'=IF(OR(C{r}="",D{r}=""),"",IF(C{r}>=100,"—",ROUND((D{r}-C{r})/(100-C{r}),2)))',
            7: f"=A_kirish!{A['daraja']}{r}",
            8: f"=B_yakuniy!{B['daraja']}{r}",
            9: (f'=IF(OR(G{r}="",H{r}=""),"",IF(MATCH(H{r},{D},0)<MATCH(G{r},{D},0),"↑",'
                f'IF(MATCH(H{r},{D},0)=MATCH(G{r},{D},0),"=","↓")))'),
            10: f'=IF(E{r}="","",IF(ABS(E{r})<=7,"o‘zgarishsiz (±2 savol)",IF(E{r}>0,"o‘sdi","pasaydi")))',
            11: f"=A_kirish!{A['tayanch']}{r}",
            12: f"=B_yakuniy!{B['tayanch']}{r}",
            13: f"=A_kirish!{A['yetti']}{r}",
            14: f"=B_yakuniy!{B['yetti']}{r}",
            15: f"=Amaliy!J{r}",
            16: f"=Amaliy!R{r}",
            17: f"=A_kirish!{A['guruh']}{r}",
        }
        for c_, v in f.items():
            so.cell(r, c_, v)
        for c_ in range(1, len(hdr) + 1):
            so.cell(r, c_).border = CHEGARA
            if 3 <= c_ <= 17:
                so.cell(r, c_).fill = HISOB_FILL
                so.cell(r, c_).alignment = Alignment(horizontal="center")
    rang_qoidalari(so, f"G{BIRINCHI_QATOR}:H{oxir}", ["Yuqori", "Yaxshi", "O‘rta", "Boshlang‘ich"])
    rang_qoidalari(so, f"I{BIRINCHI_QATOR}:I{oxir}", ["↑", "=", "↓"])
    rang_qoidalari(so, f"Q{BIRINCHI_QATOR}:Q{oxir}", ["Ilg‘or", "Asosiy", "Yordam"])
    for c_, w in zip(range(1, 20), [4, 30, 9, 9, 9, 7, 13, 13, 10, 22, 9, 9, 9, 9, 9, 9, 11, 8, 30]):
        so.column_dimensions[col(c_)].width = w
    so.row_dimensions[4].height = 30
    so.freeze_panes = "C6"

    # Sinf xulosasi (U ustundan)
    R1, R2 = BIRINCHI_QATOR, oxir

    def rng(L):
        return f"{L}{R1}:{L}{R2}"

    def jufts(x):
        return f'SUMPRODUCT(--({rng("C")}<>""),--({rng("D")}<>""),{x})'

    def g_formula(x, y):
        n_ = f'SUMPRODUCT(--({rng(x)}<>""),--({rng(y)}<>""))'
        sa = f'SUMPRODUCT(--({rng(x)}<>""),--({rng(y)}<>""),{rng(x)})'
        sb = f'SUMPRODUCT(--({rng(x)}<>""),--({rng(y)}<>""),{rng(y)})'
        return f'=IFERROR(ROUND(({sb}-{sa})/({n_}*100-{sa}),2),"")'

    U, V, W, X = 21, 22, 23, 24
    so.cell(3, U, "SINF XULOSASI").font = Font(bold=True, size=12)
    for c_, t in zip((U, V, W, X), ("Ko‘rsatkich", "Oktabr (A)", "Yanvar (B)", "O‘zgarish")):
        so.cell(4, c_, t)
    sarlavha(so, 4, (U, V, W, X))

    def ort(L):
        return f'=IFERROR(ROUND(AVERAGE({rng(L)}),1),"")'

    def ifc(L, fn):
        return f'=IF(COUNT({rng(L)})=0,"",{fn}({rng(L)}))'

    jad = [
        ("Qatnashganlar", f"=COUNT({rng('C')})", f"=COUNT({rng('D')})", True),
        ("O‘rtacha Jami %", ort("C"), ort("D"), True),
        ("Median %", ifc("C", "MEDIAN"), ifc("D", "MEDIAN"), True),
        ("Eng past %", ifc("C", "MIN"), ifc("D", "MIN"), True),
        ("Eng yuqori %", ifc("C", "MAX"), ifc("D", "MAX"), True),
        ("Tayanch % (S1–S3)", ort("K"), ort("L"), True),
        ("7-sinf % (S4–S6)", ort("M"), ort("N"), True),
        ("Amaliy %", ort("O"), ort("P"), True),
        ("Yuqori (soni)", f'=COUNTIF({rng("G")},"Yuqori")', f'=COUNTIF({rng("H")},"Yuqori")', True),
        ("Yaxshi (soni)", f'=COUNTIF({rng("G")},"Yaxshi")', f'=COUNTIF({rng("H")},"Yaxshi")', True),
        ("O‘rta (soni)", f'=COUNTIF({rng("G")},"O‘rta")', f'=COUNTIF({rng("H")},"O‘rta")', True),
        ("Boshlang‘ich (soni)", f'=COUNTIF({rng("G")},"Boshlang‘ich")', f'=COUNTIF({rng("H")},"Boshlang‘ich")', True),
        ("7-sinfga «Mos» (soni)", f'=COUNTIF(A_kirish!{A["mos"]}{R1}:{A["mos"]}{R2},"Mos")',
         f'=COUNTIF(B_yakuniy!{B["mos"]}{R1}:{B["mos"]}{R2},"Mos")', True),
    ]
    r = 5
    r_lev0 = 5 + [x[0] for x in jad].index("Yuqori (soni)")  # diagramma uchun
    for nomi, fa, fb, ozg in jad:
        so.cell(r, U, nomi)
        so.cell(r, V, fa)
        so.cell(r, W, fb)
        if ozg:
            so.cell(r, X, f'=IFERROR(IF(OR(V{r}="",W{r}=""),"",W{r}-V{r}),"")')
        r += 1
    r += 1
    so.cell(r, U, "O‘SISH (faqat ikkala testda qatnashganlar)").font = Font(bold=True)
    r += 1
    so.cell(r, U, "Juftliklar soni")
    so.cell(r, V, f'=SUMPRODUCT(--({rng("C")}<>""),--({rng("D")}<>""))')
    r += 1
    r_g = r
    for nomi, x, y in (("Sinf g — jami", "C", "D"), ("Sinf g — tayanch (S1–S3)", "K", "L"),
                       ("Sinf g — 7-sinf (S4–S6) ★ asosiy", "M", "N")):
        so.cell(r, U, nomi)
        so.cell(r, V, g_formula(x, y))
        so.cell(r, V).font = Font(bold=True)
        r += 1
    so.cell(r, U, "g talqini (jami)")
    so.cell(r, V, f'=IF(V{r_g}="","",IF(V{r_g}>=0.7,"Yuqori o‘sish",IF(V{r_g}>=0.3,"O‘rtacha o‘sish",'
                  f'IF(V{r_g}>=0,"Past o‘sish","Pasaygan"))))')
    r += 1
    for belgi, nomi in (("↑", "Daraja ko‘tarilgan"), ("=", "Daraja o‘zgarmagan"), ("↓", "Daraja tushgan")):
        so.cell(r, U, nomi)
        so.cell(r, V, f'=COUNTIF({rng("I")},"{belgi}")')
        r += 1

    r += 1
    r_bol = r
    for c_, t in zip((U, V, W, X), ("Bo‘lim", "Oktabr %", "Yanvar %", "O‘sish")):
        so.cell(r, c_, t)
    sarlavha(so, r, (U, V, W, X))
    nomlar = {x["id"]: x["nomi"] for x in a["bolimlar"]}
    for j, bid in enumerate(A["bolim_idlar"]):
        r += 1
        so.cell(r, U, f"{bid} {nomlar[bid]}")
        LA, LB, soni = A["bolimlar"][j], B["bolimlar"][j], A["bolim_soni"][j]
        so.cell(r, V, f'=IFERROR(ROUND(AVERAGE(A_kirish!{LA}{R1}:{LA}{R2})/{soni}*100,0),"")')
        so.cell(r, W, f'=IFERROR(ROUND(AVERAGE(B_yakuniy!{LB}{R1}:{LB}{R2})/{soni}*100,0),"")')
        so.cell(r, X, f'=IFERROR(IF(OR(V{r}="",W{r}=""),"",W{r}-V{r}),"")')
    r_bol_end = r
    for rr in range(5, r_bol_end + 1):
        for c_ in (U, V, W, X):
            if so.cell(rr, c_).value is not None or c_ == U:
                so.cell(rr, c_).border = CHEGARA
    so.column_dimensions[col(U)].width = 38
    for c_ in (V, W, X):
        so.column_dimensions[col(c_)].width = 12

    ch = BarChart()
    ch.type = "col"
    ch.title = "Bo‘limlar bo‘yicha sinf o‘rtachasi (%)"
    ch.y_axis.title = "%"
    ch.y_axis.scaling.min = 0
    ch.y_axis.scaling.max = 100
    ch.add_data(Reference(so, min_col=V, max_col=W, min_row=r_bol, max_row=r_bol_end), titles_from_data=True)
    ch.set_categories(Reference(so, min_col=U, min_row=r_bol + 1, max_row=r_bol_end))
    ch.height, ch.width = 8, 18
    so.add_chart(ch, f"{col(U)}{r_bol_end + 2}")

    ch2 = BarChart()
    ch2.type = "col"
    ch2.title = "Darajalar taqsimoti (o‘quvchilar soni)"
    ch2.add_data(Reference(so, min_col=V, max_col=W, min_row=r_lev0, max_row=r_lev0 + 3), titles_from_data=False)
    ch2.set_categories(Reference(so, min_col=U, min_row=r_lev0, max_row=r_lev0 + 3))
    ch2.series[0].tx = SeriesLabel(v="Oktabr (A)")
    ch2.series[1].tx = SeriesLabel(v="Yanvar (B)")
    ch2.height, ch2.width = 8, 18
    so.add_chart(ch2, f"{col(U)}{r_bol_end + 19}")

    # O‘tirish xaritasi
    xa = wb.create_sheet("Xarita")
    xa["A1"] = "7A — kompyuter xonasidagi o‘tirish xaritasi (2 hafta o‘zgarmaydi). Har katakka ism va 1 ta belgi (qiziqish) yozing."
    xa["A1"].font = Font(bold=True)
    xa.merge_cells("B3:K3")
    xa["B3"] = "DOSKA / USTOZ STOLI"
    xa["B3"].alignment = Alignment(horizontal="center")
    xa["B3"].fill = PatternFill("solid", fgColor="D9D9D9")
    xa["B3"].font = Font(bold=True)
    raqam = 1
    for qator in range(6):
        for ust in range(5):
            r0, c0 = 5 + qator * 4, 2 + ust * 2
            xa.merge_cells(start_row=r0, start_column=c0, end_row=r0, end_column=c0 + 1)
            xa.merge_cells(start_row=r0 + 1, start_column=c0, end_row=r0 + 1, end_column=c0 + 1)
            xa.merge_cells(start_row=r0 + 2, start_column=c0, end_row=r0 + 2, end_column=c0 + 1)
            xa.cell(r0, c0, f"Komp. №{raqam}").font = Font(size=9, color="666666")
            xa.cell(r0 + 1, c0).font = Font(bold=True, size=12)
            for rr in range(r0, r0 + 3):
                for cc in (c0, c0 + 1):
                    xa.cell(rr, cc).border = CHEGARA
            raqam += 1
    for c_ in range(2, 12):
        xa.column_dimensions[col(c_)].width = 12

    tartib = ["Yoriqnoma", "Royxat", "A_kirish", "Amaliy", "B_yakuniy", "Solishtirish", "Xarita", "Darajalar"]
    wb._sheets = [wb[x] for x in tartib]
    wb.active = 0
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


# ---------------------------------------------------------------- HTML / PDF

CSS = """
@page { size: A4; margin: 13mm 12mm; }
* { box-sizing: border-box; }
body { font-family: "DejaVu Sans", "Liberation Sans", Arial, sans-serif; font-size: 10pt; color: #111; line-height: 1.35; }
h1 { font-size: 15pt; margin: 0 0 6px; }
h2 { font-size: 12pt; margin: 14px 0 6px; padding: 3px 6px; background: #eef2f7; border-left: 4px solid #1f3a5f; }
h3 { font-size: 11pt; margin: 10px 0 4px; }
table { border-collapse: collapse; width: 100%; margin: 6px 0; font-size: 9pt; }
th, td { border: 1px solid #888; padding: 3px 5px; vertical-align: top; }
th { background: #1f3a5f; color: #fff; }
blockquote { margin: 6px 0; padding: 4px 10px; border-left: 3px solid #999; background: #f6f6f6; }
code { font-size: 9pt; }
ul.contains-task-list, ul.task-list { list-style: none; padding-left: 4px; }
hr { border: 0; border-top: 1px solid #bbb; margin: 10px 0; }
.qcols { column-count: 2; column-gap: 9mm; }
.sec { column-span: all; font-weight: 700; background: #eef2f7; border-left: 4px solid #1f3a5f; padding: 3px 6px; margin: 10px 0 6px; }
.q { break-inside: avoid; margin: 0 0 9px; }
.q .t { font-weight: 600; }
.q .o { margin: 2px 0 0 10px; }
.q .o span { display: inline-block; width: 11px; height: 11px; border: 1.3px solid #333; border-radius: 50%; margin-right: 5px; vertical-align: -1px; }
.fields { display: flex; gap: 18px; border: 1.5px solid #333; padding: 7px 10px; margin: 8px 0; font-size: 10.5pt; }
.fields div { flex: 1; white-space: nowrap; }
.fields div:first-child { flex: 2; }
.note { font-size: 9pt; color: #333; background: #f6f6f6; padding: 5px 8px; border-radius: 4px; }
.sheet { height: 132mm; border-bottom: 1.5px dashed #777; padding-top: 2mm; break-inside: avoid; }
.sheet:nth-child(2n) { border-bottom: 0; }
.grid { display: flex; gap: 10mm; margin-top: 6px; }
.grid .colm { flex: 1; }
.row { display: flex; align-items: center; height: 8.4mm; border-bottom: 1px solid #ddd; }
.row b { width: 9mm; text-align: right; margin-right: 4mm; }
.bub { display: inline-flex; align-items: center; justify-content: center; width: 6.2mm; height: 6.2mm; border: 1.3px solid #333; border-radius: 50%; margin-right: 2.4mm; font-size: 8pt; color: #555; }
.tent { height: 128mm; position: relative; border-bottom: 1.5px dashed #777; break-inside: avoid; }
.tent:nth-child(2n) { border-bottom: 0; }
.half { height: 50%; padding: 6mm 8mm; position: relative; }
.half.top { transform: rotate(180deg); border-bottom: 1px dotted #aaa; }
.half .name { position: absolute; left: 8mm; right: 8mm; bottom: 15mm; border-bottom: 1.5px solid #bbb; text-align: center; color: #bbb; font-size: 9pt; }
.half .lab { position: absolute; bottom: 5mm; font-size: 8pt; color: #888; }
.half .lab.l { left: 8mm; } .half .lab.r { right: 8mm; }
.half .hint { font-size: 8pt; color: #aaa; }
"""


def sahifa(sarlavha_, tana):
    return (f"<!doctype html><html lang='uz'><head><meta charset='utf-8'><title>{html.escape(sarlavha_)}</title>"
            f"<style>{CSS}</style></head><body>{tana}</body></html>")


def test_html(test):
    e = html.escape
    bolim_nomi = {b["id"]: b["nomi"] for b in test["bolimlar"]}
    q = []
    n = 0
    for i, (bid, savollar) in enumerate(bolim_guruhlari(test), 1):
        q.append(f"<div class='sec'>{i}-bo‘lim. {e(bolim_nomi[bid])}</div>")
        for s in savollar:
            n += 1
            ops = "".join(f"<div class='o'><span></span><b>{HARF[j]})</b> {e(v)}</div>" for j, v in enumerate(s["variantlar"]))
            q.append(f"<div class='q'><div class='t'>{n}. {e(s['matn'])}</div>{ops}</div>")
    tana = (
        f"<h1>7A · {e(test['nomi'])}</h1>"
        f"<div>{e(test['fan'])} · {test['vaqt_daqiqa']} daqiqa · {len(test['savollar'])} ta savol</div>"
        "<div class='fields'><div>F.I.: ____________________________</div><div>Sana: __________</div>"
        "<div>Kompyuter №: _____</div></div>"
        f"<div class='note'><b>Ko‘rsatma.</b> {e(test['korsatma'])} Javobni <b>javoblar varaqasiga</b> belgilang "
        "(yoki ustoz aytganidek shu varaqda doirachani bo‘yang).</div>"
        f"<div class='qcols'>{''.join(q)}</div>"
    )
    return sahifa(test["nomi"], tana)


def javoblar_varaqasi_html(n=30):
    def bitta():
        ust = []
        per = (n + 2) // 3
        for c in range(3):
            rows = []
            for i in range(c * per + 1, min(n, (c + 1) * per) + 1):
                bub = "".join(f"<span class='bub'>{h}</span>" for h in HARF)
                rows.append(f"<div class='row'><b>{i}.</b>{bub}</div>")
            ust.append(f"<div class='colm'>{''.join(rows)}</div>")
        return (
            "<div class='sheet'><h1>7A · Informatika · Javoblar varaqasi</h1>"
            "<div class='fields'><div>F.I.: __________________________</div><div>Sana: __________</div>"
            "<div>Variant: <span class='bub'>A</span><span class='bub'>B</span></div></div>"
            "<div class='note'>To‘g‘ri deb o‘ylagan javobingiz harfini <b>to‘liq bo‘yang</b>. Bitta savolga — bitta javob. "
            "Bilmasangiz — bo‘sh qoldiring.</div>"
            f"<div class='grid'>{''.join(ust)}</div></div>"
        )
    return sahifa("Javoblar varaqasi", bitta() + bitta())


def ism_kartochka_html(soni=2):
    def yarim(cls):
        return (f"<div class='half {cls}'><div class='hint'>Katta harflarda — qanday chaqirishimni xohlayman:</div>"
                "<div class='name'>ISM</div><div class='lab l'>Sevimli mashg‘ulotim (rasm yoki so‘z): ________</div>"
                "<div class='lab r'>Kompyuter №: ____</div></div>")
    kart = "".join(f"<div class='tent'>{yarim('top')}{yarim('bot')}</div>" for _ in range(soni))
    izoh = ("<p style='font-size:8pt;color:#777;margin:0 0 2mm'>1 varaq = 2 kartochka — kerakli nusxada chop eting. "
            "Kesib oling, nuqtali chiziqdan bukib, partaga «uycha» qilib qo‘ying. Ism ikkala tomonda yoziladi.</p>")
    return sahifa("Ism kartochkalari", izoh + kart)


def md_html(md_text, sarlavha_):
    r = subprocess.run(["pandoc", "-f", "gfm", "-t", "html5"], input=md_text.encode("utf-8"),
                       capture_output=True, check=True)
    return sahifa(sarlavha_, r.stdout.decode("utf-8"))


def chromium():
    for p in (os.environ.get("CHROMIUM"), "/opt/pw-browsers/chromium-1194/chrome-linux/chrome",
              shutil.which("chromium"), shutil.which("chromium-browser"), shutil.which("google-chrome")):
        if p and Path(p).exists():
            return p
    return None


def pdf(html_text, out, chrome):
    with tempfile.TemporaryDirectory() as td:
        src = Path(td) / "page.html"
        src.write_text(html_text, "utf-8")
        subprocess.run([chrome, "--headless=new", "--no-sandbox", "--disable-gpu", "--no-pdf-header-footer",
                        f"--print-to-pdf={out}", src.as_uri()], check=True, capture_output=True, timeout=120)
    print("  pdf:", out.relative_to(ROOT))


# ---------------------------------------------------------------- main

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--pdfsiz", action="store_true", help="PDF va DOCX yasamaslik")
    ap.add_argument("--ismlar", help="ism-familiyalar ro‘yxati (har qatorda bitta) — shaxsiy Excel uchun")
    ap.add_argument("--excel", help="Excel chiqish yo‘li (standart: jadvallar/7A_monitoring.xlsx)")
    args = ap.parse_args()

    a, b = load("test_A.json"), load("test_B.json")
    s, f = load("sorovnoma.json"), load("fikr_sorovnoma.json")
    assert len(a["savollar"]) == len(b["savollar"])
    for qa, qb in zip(a["savollar"], b["savollar"]):
        assert qa["bolim"] == qb["bolim"] and qa["qiyinlik"] == qb["qiyinlik"], (qa["id"], qb["id"])

    ismlar = None
    if args.ismlar:
        ismlar = [x.strip() for x in Path(args.ismlar).read_text("utf-8").splitlines() if x.strip()]
    excel = Path(args.excel) if args.excel else ROOT / "jadvallar" / "7A_monitoring.xlsx"
    build_excel(excel, a, b, ismlar)
    print("  xlsx:", excel)
    if args.ismlar:
        return  # shaxsiy Excel — boshqa fayllarga tegmaymiz

    md = {
        "03_oquvchi_sorovnomasi.md": sorovnoma_md(s, f),
        "04_kirish_testi_A.md": test_md(a, "04"),
        "05_yakuniy_test_B.md": test_md(b, "05"),
        "06_javoblar_kaliti.md": kalit_md(a, b),
    }
    for name, text in md.items():
        (ROOT / name).write_text(text.rstrip() + "\n", "utf-8")
        print("  md:", name)

    if args.pdfsiz:
        return
    chrome = chromium()
    if not chrome or not shutil.which("pandoc"):
        print("  ! Chromium yoki pandoc topilmadi — PDF/DOCX o‘tkazib yuborildi")
        return
    out = ROOT / "chop_etish"
    out.mkdir(exist_ok=True)
    pdf(test_html(a), out / "kirish_testi_A.pdf", chrome)
    pdf(test_html(b), out / "yakuniy_test_B.pdf", chrome)
    pdf(javoblar_varaqasi_html(len(a["savollar"])), out / "javoblar_varaqasi.pdf", chrome)
    pdf(ism_kartochka_html(), out / "ism_kartochkalari.pdf", chrome)
    pdf(md_html(sorovnoma_qism_md(s, "#"), s["nomi"]), out / "sorovnoma.pdf", chrome)
    pdf(md_html(sorovnoma_qism_md(f, "#"), f["nomi"]), out / "fikr_sorovnoma.pdf", chrome)
    for src, dst in (("01_sinf_rahbari_bilan_suhbat.md", "sinf_rahbari_suhbat.pdf"),
                     ("07_amaliy_topshiriq.md", "amaliy_topshiriq.pdf"),
                     ("08_kuzatuv_varaqasi.md", "kuzatuv_varaqasi.pdf"),
                     ("11_sinf_pasporti_shablon.md", "sinf_pasporti.pdf")):
        pdf(md_html((ROOT / src).read_text("utf-8"), src), out / dst, chrome)

    fayllar = ["README.md"] + sorted(p.name for p in ROOT.glob("[01][0-9]_*.md"))
    docx = out / "7A_toliq_qollanma.docx"
    subprocess.run(["pandoc", "-f", "gfm", "-o", str(docx), "--toc", "--metadata", "title=7A — to‘liq qo‘llanma",
                    *[str(ROOT / x) for x in fayllar]], check=True)
    print("  docx:", docx.relative_to(ROOT))


if __name__ == "__main__":
    main()
